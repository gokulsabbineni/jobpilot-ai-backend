from io import BytesIO
from pathlib import Path

import httpx
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session

from app.db import get_db
from app.models import Resume, Job
from app.models_resume_check import ResumeCheck
from app.deps import active_user
from app.config import settings
from app.services.resume_analyzer import analyze_resume
from app.api.resumes import supabase_object_url, supabase_headers


r = APIRouter(prefix="/api/user/resume-check", tags=["resume-check"])


def _serialize(x):
    return {
        "id": x.id,
        "resume_id": x.resume_id,
        "overall_score": x.overall_score,
        "ats_score": x.ats_score,
        "content_score": x.content_score,
        "impact_score": x.impact_score,
        "readability_score": x.readability_score,
        "job_alignment_score": x.job_alignment_score,
        "analysis": x.analysis,
        "updated_at": x.updated_at,
    }


async def _download_resume(resume: Resume) -> bytes:
    if not resume.file_path:
        raise HTTPException(400, "Your resume file could not be located. Please upload it again.")

    if resume.file_path.startswith("supabase://"):
        prefix = f"supabase://{settings.supabase_storage_bucket}/"
        if not resume.file_path.startswith(prefix):
            raise HTTPException(400, "Your resume storage reference is invalid.")
        object_path = resume.file_path[len(prefix):]
        try:
            async with httpx.AsyncClient(timeout=60) as client:
                response = await client.get(
                    supabase_object_url(object_path),
                    headers=supabase_headers(),
                )
                response.raise_for_status()
                return response.content
        except (httpx.HTTPError, RuntimeError) as exc:
            raise HTTPException(
                502,
                "We found your resume record, but could not read the stored file. Please try uploading the resume again.",
            ) from exc

    path = Path(resume.file_path)
    if not path.exists() or not path.is_file():
        raise HTTPException(400, "Your resume file is no longer available. Please upload it again.")
    try:
        return path.read_bytes()
    except OSError as exc:
        raise HTTPException(400, "Your resume file could not be read. Please upload it again.") from exc


def _extract_resume_text(data: bytes, file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()

    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(data))
        chunks = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))

        return "\n".join(chunks).strip()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()

    if suffix == ".txt":
        return data.decode("utf-8", errors="replace").strip()

    raise HTTPException(
        400,
        "Resume Check supports readable PDF and DOCX resumes. Please upload a readable file.",
    )


async def _ensure_resume_text(resume: Resume, db: Session) -> str:
    if resume.content_text and resume.content_text.strip():
        return resume.content_text.strip()

    data = await _download_resume(resume)
    text = _extract_resume_text(data, resume.file_name)

    if not text:
        raise HTTPException(
            400,
            "We found your resume, but could not extract readable text from it. Please upload a text-readable PDF or DOCX.",
        )

    # Cache extracted text so subsequent checks do not need to download the file again.
    resume.content_text = text
    db.commit()
    return text


@r.get("")
def get_check(db: Session = Depends(get_db), u=Depends(active_user)):
    x = db.query(ResumeCheck).filter(ResumeCheck.user_id == u.id).first()
    if not x:
        return None
    return _serialize(x)


@r.post("/analyze")
async def run_check(db: Session = Depends(get_db), u=Depends(active_user)):
    resume = db.query(Resume).filter(Resume.user_id == u.id).first()

    if not resume:
        raise HTTPException(400, "Upload a resume before running Resume Check.")

    content_text = await _ensure_resume_text(resume, db)

    prefs = u.preferences
    pref_data = {"job_titles": prefs.job_titles if prefs else []}

    # Make the analysis market-aware by comparing the resume with jobs
    # JobPilot has already discovered for the user's target role.
    target_titles = prefs.job_titles if prefs and prefs.job_titles else []
    market_jobs = []
    if target_titles:
        seen = set()
        for term in target_titles[:5]:
            rows = (
                db.query(Job)
                .filter(Job.title.ilike(f"%{term.strip()}%"))
                .order_by(Job.posted_at.desc().nullslast())
                .limit(50)
                .all()
            )
            for job in rows:
                if job.id not in seen:
                    seen.add(job.id)
                    market_jobs.append(job)
                if len(market_jobs) >= 100:
                    break
            if len(market_jobs) >= 100:
                break

    job_data = [
        {
            "title": job.title,
            "description": job.description,
            "company": job.company,
            "location": job.location,
        }
        for job in market_jobs
    ]
    analysis = analyze_resume(content_text, pref_data, job_data)

    x = db.query(ResumeCheck).filter(ResumeCheck.user_id == u.id).first()
    if not x:
        x = ResumeCheck(user_id=u.id, resume_id=resume.id)
        db.add(x)

    x.resume_id = resume.id
    s = analysis["scores"]
    x.overall_score = s["overall"]
    x.ats_score = s["ats"]
    x.content_score = s["content"]
    x.impact_score = s["impact"]
    x.readability_score = s["readability"]
    x.job_alignment_score = s["job_alignment"]
    x.analysis = analysis

    db.commit()
    db.refresh(x)
    return _serialize(x)
